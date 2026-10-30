from collections import defaultdict
from zipfile import BadZipFile

from django.contrib.auth.hashers import make_password
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import IntegrityError, transaction
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .models import Section, Student


REQUIRED_COLUMNS = {
    "اسم الطالب": "full_name",
    "البريد الإلكتروني": "email",
    "كلمة المرور": "password",
    "التخصص": "major",
    "الشعبة": "group",
    "الفترة": "period",
    "الرقم الجامعي": "academic_number",
}


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).replace("\u00a0", " ").strip()


def import_students(uploaded_file, update_existing=False):
    imported = 0
    updated = 0
    issues = []
    seen = defaultdict(set)
    workbook = None

    try:
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = [_cell_text(value) for value in next(rows, ())]
        column_indexes = {
            field_name: headers.index(header)
            for header, field_name in REQUIRED_COLUMNS.items()
            if header in headers
        }
        missing = [header for header in REQUIRED_COLUMNS if header not in headers]
        if missing:
            return 0, 0, [f"Missing required columns: {', '.join(missing)}."]

        for row_number, row in enumerate(rows, start=2):
            values = {
                field_name: _cell_text(row[index]) if index < len(row) else ""
                for field_name, index in column_indexes.items()
            }
            if not any(values.values()):
                continue

            error = _validate_row(values, seen)
            if error:
                issues.append(f"Row {row_number}: {error}")
                continue

            email = values["email"].lower()
            academic_number = values["academic_number"]
            student_by_email = Student.objects.filter(email__iexact=email).first()
            student_by_number = Student.objects.filter(academic_number=academic_number).first()
            existing_student = None
            if student_by_email or student_by_number:
                if (
                    update_existing
                    and student_by_email
                    and student_by_number
                    and student_by_email.pk == student_by_number.pk
                ):
                    existing_student = student_by_email
                else:
                    issues.append(f"Row {row_number}: email or university number already exists; row skipped.")
                    continue

            section_name = f"Group {values['group']} - {values['period']}"
            try:
                with transaction.atomic():
                    section, _ = Section.objects.get_or_create(
                        name=section_name,
                        defaults={"group_number": values["group"]},
                    )
                    student = existing_student or Student()
                    student.email = email
                    student.full_name = values["full_name"]
                    student.academic_number = academic_number
                    student.section = section
                    student.major = values["major"]
                    student.set_student_password(values["password"])
                    student.save()
            except IntegrityError:
                issues.append(f"Row {row_number}: email or university number already exists; row skipped.")
                continue
            if existing_student:
                updated += 1
            else:
                imported += 1
    except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError) as error:
        raise ValidationError("The uploaded file could not be read as a valid .xlsx workbook.") from error
    finally:
        if workbook is not None:
            workbook.close()

    return imported, updated, issues


def _validate_row(values, seen):
    required = ("full_name", "email", "password", "major", "group", "period", "academic_number")
    if any(not values[field] for field in required):
        return "one or more required student fields are blank."
    try:
        validate_email(values["email"])
    except ValidationError:
        return "email address is invalid."
    if len(values["full_name"]) > 160 or len(values["major"]) > 120:
        return "name or specialization is too long."
    if len(values["email"]) > 254:
        return "email address is too long."
    if len(values["academic_number"]) > 40:
        return "university number is too long."
    section_name = f"Group {values['group']} - {values['period']}"
    if len(section_name) > 80:
        return "group and period make a section name that is too long."
    for field in ("email", "academic_number"):
        normalized = values[field].lower() if field == "email" else values[field]
        if normalized in seen[field]:
            return "duplicate email or university number in this workbook."
        seen[field].add(normalized)
    return ""
